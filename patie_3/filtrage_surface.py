# =====================================================
# Auteur : Saleh BA-ZIGHIFAN
# Script : filtrage_surface.py
# Objectif : Filtrer les logements dont la surface réelle
#            est comprise entre 9 et 14 m² (bornes incluses)
# =====================================================

import pandas as pd
import unicodedata
from datetime import datetime
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

# ---------- Réglages ----------
DATA_PATH = "C:/Users/Saleh/Test_Saleh/data/BDD_TEST.csv"
RESULTS_DIR = "C:/Users/Saleh/Test_Saleh/patie_3/resultats"
HTML_PATH = os.path.join(RESULTS_DIR, "filtrage_resultats.html")
EXCEL_PATH = os.path.join(RESULTS_DIR, "filtrage_resultats.xlsx")
TXT_PATH = os.path.join(RESULTS_DIR, "output.txt")

BORNE_MIN = 9
BORNE_MAX = 14

# ---------- Fonctions utilitaires ----------
def to_numeric_safe(s: pd.Series) -> pd.Series:
    """Convertit proprement les valeurs de surface en numérique."""
    cleaned = (
        s.astype(str)
         .str.replace("\u00A0", " ", regex=False)
         .str.replace(",", ".", regex=False)
         .str.replace("m²", "", regex=False)
         .str.replace("m2", "", regex=False)
         .str.replace("'", "", regex=False)
         .str.replace(" ", "", regex=False)
    )
    return pd.to_numeric(cleaned, errors="coerce")

# ---------- Lecture du fichier ----------
df = pd.read_csv(DATA_PATH)

# ---------- Colonne de surface ----------
col_surface = "surface_reelle_bati"

# ---------- Conversion numérique ----------
df[col_surface] = to_numeric_safe(df[col_surface])

# ---------- Filtrage ----------
filtre = (df[col_surface] >= BORNE_MIN) & (df[col_surface] <= BORNE_MAX)
logements_filtres = df[filtre].copy()
nombre_lignes = len(logements_filtres)

# ---------- Colonnes importantes à garder ----------
colonnes_importantes = [
    "surface_reelle_bati",
    "id_mutation",
    "date_mutation",
    "nature_mutation",
    "valeur_fonciere",
    "nom_commune",
    "code_postal",
    "type_local",
    "nombre_pieces_principales",
    "surface_terrain",
    "longitude",
    "latitude",
]

colonnes_presentes = [c for c in colonnes_importantes if c in logements_filtres.columns]
df_export = logements_filtres[colonnes_presentes]

# ---------- Sauvegarde Excel ----------
os.makedirs(RESULTS_DIR, exist_ok=True)

with pd.ExcelWriter(EXCEL_PATH, engine="openpyxl") as writer:
    df_export.to_excel(writer, sheet_name="Résultats filtrés", index=False)

# ---------- Mise en forme du fichier Excel ----------
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

# Appliquer un style spécial au titre "surface_reelle_bati"
for cell in ws[1]:
    if cell.value == "surface_reelle_bati":
        cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
        # Style léger pour les autres colonnes
        cell.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

        # ---------- Ajouter des bordures à tout le tableau ----------
from openpyxl.styles import Border, Side

thin_border = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000")
)

# Appliquer les bordures à chaque cellule
for row in ws.iter_rows():
    for cell in row:
        cell.border = thin_border


wb.save(EXCEL_PATH)
wb.close()

print(f"✅ Fichier Excel créé et stylisé : {EXCEL_PATH}")

# ---------- Résumé texte ----------
with open(TXT_PATH, "w", encoding="utf-8") as f:
    f.write(f"Nombre total de logements entre {BORNE_MIN} et {BORNE_MAX} m² : {nombre_lignes}\n\n")
    f.write("Méthode utilisée :\n")
    f.write("- Lecture du fichier CSV avec pandas.\n")
    f.write(f"- Filtrage de la colonne '{col_surface}' entre {BORNE_MIN} et {BORNE_MAX} m².\n")
    f.write("- Export des lignes correspondantes dans un fichier Excel propre et lisible.\n")

print(f"✅ Résumé enregistré dans : {TXT_PATH}")

# ---------- Préparation du tableau HTML ----------
preview_rows = 50
preview_html = df_export.head(preview_rows).to_html(
    border=0, justify="center", classes="table table-hover", index=False
)

stats_surface = (
    logements_filtres[col_surface]
    .describe()
    .to_frame()
    .rename(columns={col_surface: "surface (m²)"})
    .to_html(border=0, justify="center", classes="table table-hover")
)

# ---------- Génération du rapport HTML ----------
date_gen = datetime.now().strftime("%d/%m/%Y à %H:%M")

with open(HTML_PATH, "w", encoding="utf-8") as f:
    f.write(f"""
    <html>
    <head>
        <meta charset='utf-8'>
        <title>Filtrage - Surface réelle (9 à 14 m²)</title>
        <style>
            body {{
                font-family: 'Segoe UI', Tahoma, sans-serif;
                margin: 40px;
                background-color: #f4f6f8;
                color: #333;
            }}
            .card {{
                background: white; padding: 20px 24px; border-radius: 12px;
                box-shadow: 0 4px 16px rgba(0,0,0,0.08); margin: 0 auto 24px; width: 90%;
            }}
            h1 {{ text-align: center; color: #0d6efd; margin-bottom: 10px; }}
            h2 {{
                color: #0d6efd; border-left: 4px solid #0d6efd; padding-left: 8px;
                margin: 24px 0 12px;
            }}
            .kpi {{
                display: flex; gap: 18px; justify-content: center; flex-wrap: wrap; margin: 12px 0 6px;
            }}
            .kpi .item {{
                background: #e9f2ff; border-radius: 10px; padding: 10px 14px; min-width: 220px; text-align: center;
                border: 1px solid #cfe2ff;
            }}
            .muted {{ color: #666; font-size: 14px; }}
            table {{
                border-collapse: collapse; margin: auto; width: 95%; font-size: 15px; background: white;
                box-shadow: 0 2px 10px rgba(0,0,0,0.08);
            }}
            th, td {{ border: 1px solid #ddd; padding: 8px 10px; text-align: center; }}
            th {{ background-color: #0d6efd; color: white; font-weight: 600; }}
            tr:nth-child(even) {{ background-color: #f9fbff; }}
            footer {{ text-align: center; margin-top: 24px; color: #555; font-size: 14px; }}
            .note {{ text-align: center; margin: 8px 0 22px; color: #444; }}
        </style>
    </head>
    <body>
        <div class="card">
            <h1>🧮 Filtrage — « surface réelle » entre {BORNE_MIN} et {BORNE_MAX} m²</h1>
            <p class="note">Généré le {date_gen} — Colonne utilisée : <b>{col_surface}</b></p>

            <div class="kpi">
                <div class="item"><b>Nombre total de lignes filtrées</b><br><span style="font-size:22px">{nombre_lignes}</span></div>
                <div class="item"><b>Bornes de filtrage (inclusives)</b><br>{BORNE_MIN} m² → {BORNE_MAX} m²</div>
                <div class="item"><b>Export Excel</b><br><span class="muted">{EXCEL_PATH}</span></div>
            </div>

            <h2>Statistiques sur la surface (lignes filtrées)</h2>
            {stats_surface}

            <h2>Aperçu des {preview_rows} premières lignes</h2>
            {preview_html}
            <p class="muted" style="text-align:center">Aperçu limité à {preview_rows} lignes pour la lisibilité.</p>
        </div>
        <footer>
            ✅ Rapport généré automatiquement par <b>Saleh BA-ZIGHIFAN</b> — Analyse de données professionnelle
        </footer>
    </body>
    </html>
    """)

print(f"✅ Rapport HTML créé : {HTML_PATH}")
print("\n🎯 Tout est terminé avec succès !")
